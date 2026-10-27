"""File parsers by type. Each turns raw bytes into a ParsedDocument; options come from parse-phase rules."""
import csv
import datetime as dt
import email
import io
import re
import zipfile
from collections import Counter
from email import policy
from urllib.parse import urljoin

from bs4 import BeautifulSoup

from apps.core.bangla import clean_text, to_ascii_digits
from apps.core.merge import get_path
from apps.plugins.registry import register, registry

from . import ocr
from .base import ParsedDocument, ParserPlugin
from .textdecode import decode_bytes

NOISE_TAGS = ["script", "style", "noscript", "template", "svg"]


def parser_for(file_type):
    for cls in registry.all("parser").values():
        if file_type in cls.file_types:
            return cls
    return None


def html_to_text(html_or_soup):
    soup = html_or_soup if isinstance(html_or_soup, BeautifulSoup) else BeautifulSoup(html_or_soup or "", "lxml")
    for tag in soup(NOISE_TAGS):
        tag.decompose()
    for br in soup.find_all("br"):
        br.replace_with("\n")
    return clean_text(soup.get_text("\n"))


def _encoding_override(settings):
    return get_path(settings or {}, "content.encoding", "") or ""


def _apply_line_filters(text, patterns):
    if not patterns:
        return text
    compiled = [re.compile(p) for p in patterns]
    return "\n".join(line for line in text.splitlines() if not any(c.search(line) for c in compiled))


# --- HTML ----------------------------------------------------------------------------
@register
class HtmlParser(ParserPlugin):
    name = "html"
    label = "HTML page"
    file_types = ("html",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        doc = ParsedDocument(file_type="html")
        html, doc.encoding = decode_bytes(data, override=_encoding_override(settings), content_type=content_type,
                                          warnings=doc.warnings)
        soup = BeautifulSoup(html, "lxml")
        doc.html = html
        if soup.title:
            doc.title = clean_text(soup.title.get_text(" "))
        for meta in soup.find_all("meta"):
            key = meta.get("name") or meta.get("property") or meta.get("http-equiv")
            if key and meta.get("content"):
                doc.metadata[key] = meta["content"]
        if not doc.title:
            doc.title = doc.metadata.get("og:title", "")
        for a in soup.select("a[href]"):
            href = a["href"].strip()
            if href and not href.startswith(("javascript:", "mailto:", "#")):
                doc.links.append({"url": urljoin(url, href) if url else href, "text": clean_text(a.get_text(" "))})
        doc.text = html_to_text(soup)
        return doc


# --- PDF -----------------------------------------------------------------------------
def parse_page_range(spec, count):
    """'1-3,5' -> [0,1,2,4]; '-2' -> last two pages; blank -> all."""
    if not spec or not spec.strip():
        return list(range(count))
    pages = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if part.startswith("-") and part[1:].isdigit():
            pages.extend(range(max(0, count - int(part[1:])), count))
        elif "-" in part:
            start, end = (int(x) for x in part.split("-", 1))
            pages.extend(range(start - 1, min(end, count)))
        else:
            pages.append(int(part) - 1)
    seen = set()
    return [p for p in pages if 0 <= p < count and not (p in seen or seen.add(p))]


HEADER_FOOTER_LINES = 3


def _line_key(line):
    return re.sub(r"\d+", "#", to_ascii_digits(line.strip()))


def _edge_lines(text):
    lines = [line for line in text.splitlines() if line.strip()]
    return lines[:HEADER_FOOTER_LINES] + lines[-HEADER_FOOTER_LINES:]


def _repeating_lines(page_texts, threshold=0.6):
    """Short lines found in the header/footer zone of most pages (digits ignored, so 'Page 3' repeats)."""
    if len(page_texts) < 3:
        return set()
    counter = Counter()
    for text in page_texts:
        counter.update({_line_key(line) for line in _edge_lines(text) if len(line.strip()) < 120})
    limit = max(2, int(len(page_texts) * threshold))
    return {key for key, n in counter.items() if n >= limit}


def _strip_edges(text, repeated):
    lines = text.splitlines()
    content = [i for i, line in enumerate(lines) if line.strip()]
    edge = set(content[:HEADER_FOOTER_LINES] + content[-HEADER_FOOTER_LINES:])
    kept = "\n".join(line for i, line in enumerate(lines) if not (i in edge and _line_key(line) in repeated))
    return kept if kept.strip() else text  # never blank a page entirely


def _table_text(table):
    rows = []
    for row in table.extract():
        cells = [clean_text(str(c)) if c is not None else "" for c in row]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


@register
class PdfParser(ParserPlugin):
    name = "pdf"
    label = "PDF"
    file_types = ("pdf",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        import fitz

        from .rules import PdfRule

        opts = PdfRule.with_defaults(options or {})
        doc = ParsedDocument(file_type="pdf")
        pdf = fitz.open(stream=data, filetype="pdf")
        if pdf.needs_pass:
            doc.attention.append("PDF is password protected")
            return doc
        doc.pages = pdf.page_count
        doc.metadata = {f"pdf.{k}": v for k, v in (pdf.metadata or {}).items() if v}
        meta_title = (pdf.metadata or {}).get("title", "").strip()
        if meta_title and not re.fullmatch(r"(untitled|microsoft word.*|document\d*|.*\.docx?)", meta_title, re.I):
            doc.title = meta_title
        fonts, texts = set(), []
        mode, layout = opts["text_mode"], opts["layout"]
        scanned = 0
        for index in parse_page_range(opts.get("page_range", ""), pdf.page_count):
            page = pdf[index]
            fonts.update(f[3] for f in page.get_fonts() if f[3])
            text = self._page_text(page, layout)
            needs_ocr = mode == "ocr" or (mode == "text_then_ocr" and len(text.strip()) < opts["min_chars_per_page"])
            if needs_ocr:
                scanned += 1
                if ocr.ocr_available():
                    text = ocr.ocr_pdf_page(page)
                    doc.flag("ocr_used")
                elif mode == "ocr" or not text.strip():
                    doc.flag("ocr_unavailable")
            texts.append(text)
        if opts.get("remove_repeating_lines"):
            repeated = _repeating_lines(texts)
            if repeated:
                texts = [_strip_edges(t, repeated) for t in texts]
        text = "\n\n".join(t.strip() for t in texts if t.strip())
        text = _apply_line_filters(text, opts.get("remove_line_patterns"))
        doc.text = clean_text(text)
        doc.font_names = sorted(fonts)
        if "ocr_unavailable" in doc.quality_flags:
            doc.attention.append(
                f"{scanned} page(s) have no text layer (scanned) and OCR is not available: {ocr.ocr_status_message()}"
            )
        for link in (lnk for page in pdf for lnk in page.get_links() if lnk.get("uri")):
            doc.links.append({"url": link["uri"], "text": ""})
        return doc

    @staticmethod
    def _page_text(page, layout):
        if layout == "plain":
            return page.get_text("text")
        if layout == "preserve_layout":
            return page.get_text("text", sort=True)
        tables = list(page.find_tables())
        if not tables:
            return page.get_text("text", sort=True)
        boxes = [t.bbox for t in tables]
        outside = []
        for block in page.get_text("blocks", sort=True):
            x0, y0, x1, y1, text = block[:5]
            if not any(x0 >= b[0] - 1 and y0 >= b[1] - 1 and x1 <= b[2] + 1 and y1 <= b[3] + 1 for b in boxes):
                outside.append((y0, text))
        outside.extend((t.bbox[1], _table_text(t)) for t in tables)
        return "\n".join(text for _y, text in sorted(outside, key=lambda pair: pair[0]))


# --- Word ----------------------------------------------------------------------------
@register
class DocxParser(ParserPlugin):
    name = "docx"
    label = "Word .docx"
    file_types = ("docx",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        from .rules import OfficeRule

        opts = OfficeRule.with_defaults(options or {})
        include = opts["docx_include"]
        document = docx.Document(io.BytesIO(data))
        doc = ParsedDocument(file_type="docx")
        parts, fonts, first_heading = [], set(), ""
        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p" and include in {"both", "paragraphs"}:
                para = Paragraph(child, document)
                text = para.text.strip()
                if text:
                    parts.append(text)
                    if not first_heading and para.style is not None and para.style.name.lower().startswith(("heading", "title")):
                        first_heading = text
                fonts.update(run.font.name for run in para.runs if run.font.name)
                fonts.update(child.xpath(".//w:rFonts/@w:ascii | .//w:rFonts/@w:cs"))
            elif tag == "tbl" and include in {"both", "tables"}:
                table = Table(child, document)
                for row in table.rows:
                    cells, last = [], None
                    for cell in row.cells:
                        if cell._tc is not last:  # merged cells repeat the same element
                            cells.append(clean_text(cell.text))
                        last = cell._tc
                    if any(cells):
                        parts.append(" | ".join(cells))
                fonts.update(child.xpath(".//w:rFonts/@w:ascii | .//w:rFonts/@w:cs"))
        try:
            for style in document.styles:
                if getattr(style, "font", None) is not None and style.font.name:
                    fonts.add(style.font.name)
        except Exception:  # noqa: BLE001 - style quirks must not fail the parse
            pass
        core = document.core_properties
        doc.title = (core.title or "").strip() or first_heading
        doc.metadata = {k: str(v) for k, v in {"docx.title": core.title, "docx.author": core.author,
                                               "docx.created": core.created, "docx.modified": core.modified}.items() if v}
        doc.text = clean_text("\n".join(parts))
        doc.font_names = sorted(f for f in fonts if f)
        return doc


@register
class DocParser(ParserPlugin):
    name = "doc"
    label = "Word 97-2003 .doc"
    file_types = ("doc",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        doc = ParsedDocument(file_type="doc")
        doc.flag("unparsed_file_type")
        doc.attention.append("Word 97-2003 (.doc) files cannot be read automatically. The original is archived; "
                             "convert it to .docx or PDF and upload it, or paste the text.")
        return doc


# --- Spreadsheets --------------------------------------------------------------------
def _col_index(letters):
    value = 0
    for ch in letters.upper():
        value = value * 26 + (ord(ch) - 64)
    return value - 1


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return clean_text(str(value))


@register
class SpreadsheetParser(ParserPlugin):
    name = "spreadsheet"
    label = "Excel / CSV"
    file_types = ("xlsx", "xls", "csv")

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        from .rules import OfficeRule

        file_type = self.file_type_hint or "xlsx"

        opts = OfficeRule.with_defaults(options or {})
        doc = ParsedDocument(file_type=file_type)
        sheets = self._sheets(data, file_type, settings, doc)
        wanted = [s.strip() for s in opts.get("sheets") or []]
        if wanted:
            sheets = [(name, rows) for i, (name, rows) in enumerate(sheets, start=1)
                      if name in wanted or str(i) in wanted]
        row_start, row_end = 1, None
        if opts.get("row_range"):
            a, _, b = opts["row_range"].partition("-")
            row_start, row_end = int(a), (int(b) if b.strip() else None)
        col_slice = slice(None)
        if opts.get("column_range"):
            a, b = (x.strip() for x in opts["column_range"].split("-"))
            col_slice = slice(_col_index(a), _col_index(b) + 1)
        header_row = opts["header_row"]
        lines = []
        for sheet_name, rows in sheets:
            rows = [[_cell(c) for c in row][col_slice] for row in rows]
            header = rows[header_row - 1] if header_row and len(rows) >= header_row else []
            if len(sheets) > 1:
                lines.append(f"## {sheet_name}")
            for number, row in enumerate(rows, start=1):
                if number == header_row or number < row_start or (row_end and number > row_end) or not any(row):
                    continue
                if opts["row_format"] == "key_value" and header:
                    pairs = [(header[i] if i < len(header) and header[i] else f"Col{i + 1}", v) for i, v in enumerate(row) if v]
                    text = "; ".join(f"{k}: {v}" for k, v in pairs)
                    fields = dict(pairs)
                else:
                    text = "\t".join(row)
                    fields = {}
                lines.append(text)
                if opts.get("rows_as_items"):
                    doc.segments.append({"title": f"{sheet_name} row {number}", "text": text, "fields": fields})
        doc.text = "\n".join(lines)  # cells are cleaned individually; keep tabs between them
        doc.title = sheets[0][0] if len(sheets) == 1 and file_type != "csv" else ""
        return doc

    @staticmethod
    def _sheets(data, file_type, settings, doc):
        if file_type == "xlsx":
            import openpyxl

            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in book.worksheets]
        if file_type == "xls":
            import xlrd

            book = xlrd.open_workbook(file_contents=data)
            return [(s.name, [s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]
        text, doc.encoding = decode_bytes(data, override=_encoding_override(settings), warnings=doc.warnings)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [("CSV", list(csv.reader(io.StringIO(text), dialect)))]


# --- Text / RTF / images -------------------------------------------------------------
@register
class TxtParser(ParserPlugin):
    name = "txt"
    label = "Plain text"
    file_types = ("txt",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        doc = ParsedDocument(file_type="txt")
        text, doc.encoding = decode_bytes(data, override=_encoding_override(settings), content_type=content_type,
                                          warnings=doc.warnings)
        doc.text = clean_text(text)
        return doc


_RTF_FONT = re.compile(r"\\f\d+(?:\\[a-z]+-?\d*\s?)*\s*([^;{}\\]+);")


@register
class RtfParser(ParserPlugin):
    name = "rtf"
    label = "RTF"
    file_types = ("rtf",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        from striprtf.striprtf import rtf_to_text

        doc = ParsedDocument(file_type="rtf")
        raw, doc.encoding = decode_bytes(data, override=_encoding_override(settings) or "cp1252")
        doc.font_names = sorted({m.strip() for m in _RTF_FONT.findall(raw[:20000])})
        doc.text = clean_text(rtf_to_text(raw, errors="ignore"))
        return doc


@register
class ImageParser(ParserPlugin):
    name = "image"
    label = "Image (OCR)"
    file_types = ("image",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        doc = ParsedDocument(file_type="image")
        if ocr.ocr_available():
            doc.text = clean_text(ocr.ocr_image_bytes(data))
            doc.flag("ocr_used")
        else:
            doc.flag("ocr_unavailable")
            doc.attention.append(f"Image needs OCR to read: {ocr.ocr_status_message()}")
        return doc


# --- Email ---------------------------------------------------------------------------
_QUOTE_HEADERS = [
    re.compile(r"^On .{5,200} wrote:\s*$"),
    re.compile(r"^-{2,}\s*Original Message\s*-{2,}", re.I),
    re.compile(r"^From: .+$"),  # Outlook-style quoted header block
    re.compile(r"^-{2,}\s*Forwarded message\s*-{2,}", re.I),
]


def strip_email_text(text, *, strip_quoted=True, signature_patterns=(), remove_patterns=()):
    lines = text.splitlines()
    out = []
    signatures = [re.compile(p) for p in signature_patterns or []]
    for index, line in enumerate(lines):
        stripped = line.strip()
        if strip_quoted:
            if stripped.startswith(">"):
                continue
            # a quote header only counts after some real content (forwarded bulletins keep their body)
            if out and any(p.match(stripped) for p in _QUOTE_HEADERS[:2]):
                break
        if out and any(p.search(stripped) for p in signatures):
            break
        out.append(line)
    return _apply_line_filters("\n".join(out), remove_patterns)


def _part_text(part):
    try:
        return part.get_content()
    except (LookupError, UnicodeDecodeError):
        payload = part.get_payload(decode=True) or b""
        return decode_bytes(payload)[0]


@register
class EmlParser(ParserPlugin):
    name = "eml"
    label = "Email .eml"
    file_types = ("eml",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        from .rules import EmailRule

        opts = EmailRule.with_defaults(options or {})
        msg = email.message_from_bytes(data, policy=policy.default)
        doc = ParsedDocument(file_type="eml")
        headers = {}
        for key in ("subject", "from", "to", "cc", "date", "message-id"):
            try:
                value = msg.get(key)
                headers[key] = clean_text(str(value)) if value is not None else ""
            except Exception:  # noqa: BLE001 - malformed headers must not fail the parse
                headers[key] = ""
        doc.title = headers["subject"]
        doc.metadata = {f"email.{k.replace('-', '_')}": v for k, v in headers.items() if v}
        plain = msg.get_body(preferencelist=("plain",))
        html = msg.get_body(preferencelist=("html",))
        source = opts["body_source"]
        chosen_html = None
        if source == "plain" or (source == "prefer_plain" and plain is not None):
            body = _part_text(plain) if plain is not None else ""
        elif html is not None:
            chosen_html = _part_text(html)
            body = html_to_text(chosen_html)
        else:
            body = _part_text(plain) if plain is not None else ""
        if chosen_html:
            doc.html = chosen_html
        doc.text = clean_text(strip_email_text(
            body, strip_quoted=opts["strip_quoted"], signature_patterns=opts["signature_patterns"],
            remove_patterns=opts.get("remove_line_patterns"),
        ))
        for part in msg.iter_attachments():
            payload = part.get_payload(decode=True)
            if payload is None:
                continue
            disposition = (part.get_content_disposition() or "").lower()
            doc.attachments.append({
                "filename": part.get_filename() or "",
                "content_type": part.get_content_type(),
                "data": payload,
                "inline": disposition == "inline" or (bool(part.get("Content-ID")) and disposition != "attachment"),
            })
        return doc


# --- ZIP -----------------------------------------------------------------------------
MAX_ZIP_MEMBERS = 200
MAX_ZIP_TOTAL = 200 * 1024 * 1024


@register
class ZipParser(ParserPlugin):
    name = "zip"
    label = "ZIP archive"
    file_types = ("zip",)

    def parse(self, data, *, filename="", content_type="", url="", options=None, settings=None):
        doc = ParsedDocument(file_type="zip")
        max_member = float(get_path(settings or {}, "files.max_size_mb", 50)) * 1024 * 1024
        archive = zipfile.ZipFile(io.BytesIO(data))
        total, names = 0, []
        for info in archive.infolist():
            if info.is_dir() or info.filename.startswith("__MACOSX/") or info.filename.endswith("/.DS_Store"):
                continue
            if len(doc.attachments) >= MAX_ZIP_MEMBERS:
                doc.warnings.append(f"Only the first {MAX_ZIP_MEMBERS} files were taken from the ZIP.")
                break
            if info.file_size > max_member or total + info.file_size > MAX_ZIP_TOTAL:
                doc.warnings.append(f"Skipped {info.filename}: too large ({info.file_size} bytes).")
                continue
            total += info.file_size
            names.append(info.filename)
            doc.attachments.append({"filename": info.filename, "content_type": "", "data": archive.read(info)})
        doc.text = "Archive contents:\n" + "\n".join(names)
        return doc
