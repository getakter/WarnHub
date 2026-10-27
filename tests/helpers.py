"""Test-only plugins and file builders."""
import io
import zipfile
from email.message import EmailMessage

from apps.acquisition.base import Acquired, FetcherPlugin
from apps.extraction.base import ExtractorPlugin
from apps.plugins.registry import registry

STATIC_DOCUMENTS = []  # filled by tests; yielded by the test_static fetcher


class StaticFetcher(FetcherPlugin):
    name = "test_static"
    label = "Test: static documents"
    requires_entry_urls = False

    def collect(self, ctx):
        for doc in list(STATIC_DOCUMENTS):
            yield Acquired(**doc) if isinstance(doc, dict) else doc


class NoMatchRule(ExtractorPlugin):
    name = "test_no_match"
    label = "Test: never matches"

    def apply(self, items, ctx):
        ctx.no_match("nothing to see")
        return items


class ExplodingRule(ExtractorPlugin):
    name = "test_explode"
    label = "Test: raises"

    def apply(self, items, ctx):
        raise ValueError("boom")


class SplitLinesRule(ExtractorPlugin):
    """Each non-empty line becomes an item keyed by its first word."""

    name = "test_split_lines"
    label = "Test: split lines"
    produces_items = True

    def apply(self, items, ctx):
        from apps.extraction.base import Item

        out = []
        for item in items:
            for line in item.body_text.splitlines():
                if line.strip():
                    out.append(Item(title=line[:40], body_text=line, source_url=item.source_url,
                                    source_item_key=f"{item.source_url}#{line.split()[0]}"))
        return out


for _cls in (StaticFetcher, NoMatchRule, ExplodingRule, SplitLinesRule):
    registry.register(_cls)


def make_pdf(pages):
    import fitz

    pdf = fitz.open()
    for text in pages:
        page = pdf.new_page()
        y = 72
        for line in text.splitlines():
            page.insert_text((72, y), line)
            y += 16
    return pdf.tobytes()


def make_docx(paragraphs=(), table=None, heading=None):
    import docx

    document = docx.Document()
    if heading:
        document.add_heading(heading, 1)
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.rows[r].cells[c].text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets):
    import openpyxl

    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def make_zip(files):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def make_eml(subject, body, sender="ffwc@ffwc.gov.bd", to="alerts@example.org", attachments=(), html=None,
             message_id="<test-1@ffwc.gov.bd>", inline=()):
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = to
    msg["Message-ID"] = message_id
    msg["Date"] = "Tue, 27 Oct 2026 09:00:00 +0600"
    msg.set_content(body)
    if html:
        msg.add_alternative(html, subtype="html")
    for filename, data, maintype, subtype in attachments:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename)
    for filename, data, maintype, subtype in inline:
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=filename, disposition="inline",
                           cid="<logo@example.org>")
    return msg.as_bytes()


# --- Fake mailbox (protocol plugin "test_fake") -----------------------------------------
from apps.mailbox.protocols import MailboxProtocol, password_for  # noqa: E402


class FakeServer:
    """In-memory IMAP-like store shared by all fake connections."""

    folders = {}
    uidvalidity = "1"
    actions = []

    @classmethod
    def reset(cls, **folders):
        cls.folders = {name: dict(messages) for name, messages in folders.items()} or {"INBOX": {}}
        cls.actions = []
        cls.uidvalidity = "1"

    @classmethod
    def deliver(cls, folder, uid, raw):
        cls.folders.setdefault(folder, {})[str(uid)] = raw


class FakeConnection:
    def folders(self):
        return list(FakeServer.folders)

    def uidvalidity(self, folder):
        return FakeServer.uidvalidity

    def counts(self, folder):
        return {"messages": len(FakeServer.folders.get(folder, {})), "unseen": 0}

    def search_uids(self, folder, since=None, unseen_only=False):
        return sorted(FakeServer.folders.get(folder, {}), key=int)

    def fetch_raw(self, folder, uids):
        store = FakeServer.folders.get(folder, {})
        return {uid: store[uid] for uid in uids if uid in store}

    def mark_read(self, folder, uid):
        FakeServer.actions.append(("mark_read", folder, uid))

    def move(self, folder, uid, target):
        FakeServer.actions.append(("move", folder, uid, target))

    def label(self, folder, uid, label):
        FakeServer.actions.append(("label", folder, uid, label))

    def close(self):
        pass


class FakeProtocol(MailboxProtocol):
    name = "test_fake"
    label = "Test: in-memory mailbox"

    def connect(self, account):
        password_for(account)  # same env-var check as IMAP
        return FakeConnection()


registry.register(FakeProtocol)
